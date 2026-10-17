"""
Real-time report writer for CSV, XLSX (with openpyxl), and progress tracking.
Guarantees immediate persistence to disk after each analyzed PR.
"""

from __future__ import annotations

import csv
from datetime import datetime, timezone
import logging
import os
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import AuditRecord, AuditStatus, REPORT_COLUMNS


logger = logging.getLogger("git_audit.report_writer")


class ReportWriter:
    """Manages real-time streaming writes to CSV, XLSX, and progress.txt."""

    # Color definitions for styling Excel outputs
    COLOR_HEADER_FILL = "1F4E79"       # Deep Navy Blue
    COLOR_HEADER_FONT = "FFFFFF"       # White text
    COLOR_PASS_FILL = "D9EAD3"         # Soft Green
    COLOR_PASS_FONT = "274E13"         # Dark Green
    COLOR_REVIEW_FILL = "FFF2CC"       # Soft Yellow
    COLOR_REVIEW_FONT = "7F6000"       # Dark Gold/Brown
    COLOR_ERROR_FILL = "F4CCCC"        # Soft Red
    COLOR_ERROR_FONT = "990000"        # Dark Red
    COLOR_BORDER = "D9D9D9"            # Subtle light gray border

    def __init__(
        self,
        csv_path: Path,
        xlsx_path: Path,
        progress_path: Path,
        realtime_save: bool = True,
        status_colors: bool = True,
    ) -> None:
        self.csv_path = Path(csv_path)
        self.xlsx_path = Path(xlsx_path)
        self.progress_path = Path(progress_path)
        self.realtime_save = realtime_save
        self.status_colors = status_colors

        # Ensure parent directories exist
        for p in (self.csv_path, self.xlsx_path, self.progress_path):
            p.parent.mkdir(parents=True, exist_ok=True)

        self._csv_file = None
        self._csv_writer = None
        self._wb: Optional[openpyxl.Workbook] = None
        self._ws: Optional[openpyxl.worksheet.worksheet.Worksheet] = None

        self._init_csv()
        self._init_xlsx()

    def _init_csv(self) -> None:
        """Initialize CSV report file with header if not present."""
        file_exists = self.csv_path.exists() and self.csv_path.stat().st_size > 0
        self._csv_file = open(self.csv_path, "a", newline="", encoding="utf-8")
        self._csv_writer = csv.writer(self._csv_file)

        if not file_exists:
            logger.info("Initializing CSV report: %s", self.csv_path)
            self._csv_writer.writerow(REPORT_COLUMNS)
            self._csv_file.flush()
            os.fsync(self._csv_file.fileno())

    def _init_xlsx(self) -> None:
        """Initialize XLSX workbook with styled header row and freeze panes."""
        if self.xlsx_path.exists():
            try:
                self._wb = openpyxl.load_workbook(self.xlsx_path)
                self._ws = self._wb.active
                logger.info("Loaded existing XLSX report: %s", self.xlsx_path)
                return
            except Exception as ex:
                logger.warning("Could not load existing XLSX, recreating: %s", ex)

        logger.info("Creating new styled XLSX report: %s", self.xlsx_path)
        self._wb = openpyxl.Workbook()
        self._ws = self._wb.active
        self._ws.title = "Audit Report"

        # Freeze header row
        self._ws.freeze_panes = "A2"

        # Header formatting
        header_fill = PatternFill(start_color=self.COLOR_HEADER_FILL, end_color=self.COLOR_HEADER_FILL, fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color=self.COLOR_HEADER_FONT)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color=self.COLOR_BORDER),
            right=Side(style="thin", color=self.COLOR_BORDER),
            top=Side(style="thin", color=self.COLOR_BORDER),
            bottom=Side(style="thin", color=self.COLOR_BORDER),
        )

        self._ws.append(REPORT_COLUMNS)
        self._ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(REPORT_COLUMNS) + 1):
            cell = self._ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        if self.realtime_save:
            self._wb.save(self.xlsx_path)

    def write_record(self, record: AuditRecord, total_scanned: int = 0, current_repo: str = "") -> None:
        """
        Appends single audit record to CSV and XLSX, flushes buffers, and updates progress.txt.
        """
        row_values = record.to_row()

        # 1. Realtime CSV write
        try:
            self._csv_writer.writerow(row_values)
            self._csv_file.flush()
            os.fsync(self._csv_file.fileno())
            logger.debug("Flushed CSV row for PR #%s", record.pr_number)
        except Exception as ex:
            logger.error("Failed to write CSV row: %s", ex)

        # 2. Realtime XLSX write
        try:
            if self._wb and self._ws:
                self._ws.append(row_values)
                curr_row = self._ws.max_row
                self._ws.row_dimensions[curr_row].height = 20

                # Style data cells
                thin_border = Border(
                    left=Side(style="thin", color=self.COLOR_BORDER),
                    right=Side(style="thin", color=self.COLOR_BORDER),
                    top=Side(style="thin", color=self.COLOR_BORDER),
                    bottom=Side(style="thin", color=self.COLOR_BORDER),
                )
                data_font = Font(name="Segoe UI", size=10)

                for col_idx, val in enumerate(row_values, start=1):
                    cell = self._ws.cell(row=curr_row, column=col_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")

                    # Column 5: PRUrl hyperlink
                    if col_idx == 5 and val and str(val).startswith("http"):
                        cell.hyperlink = str(val)
                        cell.font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")

                    # Column 22: Status formatting
                    if col_idx == 22 and self.status_colors:
                        status_str = str(val).upper().strip()
                        if status_str == AuditStatus.PASS.value:
                            cell.fill = PatternFill(start_color=self.COLOR_PASS_FILL, end_color=self.COLOR_PASS_FILL, fill_type="solid")
                            cell.font = Font(name="Segoe UI", size=10, bold=True, color=self.COLOR_PASS_FONT)
                        elif status_str == AuditStatus.REVIEW.value:
                            cell.fill = PatternFill(start_color=self.COLOR_REVIEW_FILL, end_color=self.COLOR_REVIEW_FILL, fill_type="solid")
                            cell.font = Font(name="Segoe UI", size=10, bold=True, color=self.COLOR_REVIEW_FONT)
                        elif status_str == AuditStatus.ERROR.value:
                            cell.fill = PatternFill(start_color=self.COLOR_ERROR_FILL, end_color=self.COLOR_ERROR_FILL, fill_type="solid")
                            cell.font = Font(name="Segoe UI", size=10, bold=True, color=self.COLOR_ERROR_FONT)

                # Set AutoFilter
                self._ws.auto_filter.ref = self._ws.dimensions

                # Autosize columns
                self._adjust_column_widths()

                if self.realtime_save:
                    self._wb.save(self.xlsx_path)
                    logger.debug("Saved XLSX workbook after PR #%s", record.pr_number)

        except Exception as ex:
            logger.error("Failed to write XLSX row: %s", ex)

        # 3. Realtime progress.txt update
        try:
            self._update_progress(
                f"[{datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}] "
                f"Repo: {record.repository} | PR #{record.pr_number} | Branch: {record.branch} | "
                f"Status: {record.status} | Action: {record.action}"
            )
        except Exception as ex:
            logger.warning("Failed to update progress file: %s", ex)

    def _adjust_column_widths(self) -> None:
        """Autosize openpyxl worksheet columns based on content length."""
        if not self._ws:
            return
        for col in self._ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = str(val)
                    # Handle newlines
                    lines = val_str.split("\n")
                    longest = max(len(line) for line in lines)
                    if longest > max_len:
                        max_len = longest
            # Add padding and cap max width
            adjusted_width = min(max(max_len + 3, 12), 65)
            self._ws.column_dimensions[col_letter].width = adjusted_width

    def _update_progress(self, message: str) -> None:
        """Append line to progress.txt and flush to disk."""
        with open(self.progress_path, "a", encoding="utf-8") as pf:
            pf.write(f"{message}\n")
            pf.flush()
            os.fsync(pf.fileno())

    def update_status_message(self, message: str) -> None:
        """General progress checkpoint update."""
        ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        formatted = f"[{ts}] {message}"
        self._update_progress(formatted)
        logger.info("Progress: %s", message)

    def close(self) -> None:
        """Flush and close all open file handles safely."""
        logger.info("Finalizing and closing report writers...")
        if self._csv_file and not self._csv_file.closed:
            try:
                self._csv_file.flush()
                os.fsync(self._csv_file.fileno())
                self._csv_file.close()
            except Exception as ex:
                logger.warning("Error closing CSV file: %s", ex)

        if self._wb:
            try:
                if self._ws:
                    self._adjust_column_widths()
                self._wb.save(self.xlsx_path)
                self._wb.close()
            except Exception as ex:
                logger.warning("Error saving final XLSX workbook: %s", ex)
